"""Assemble a DOCX from a prepared title and evidence-backed content; no canned prose."""
import argparse
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from report_formatting import field, install_styles, install_toc_styles, normalize_dashes
from report_plan import load_plan, structure_issues

def build(title, content, plan, output, root=None):
    root = root or Path.cwd()
    doc = Document(title)
    if len(doc.sections) != 1:
        raise ValueError('Prepared title must have one section; remove sample body and trailing section breaks.')
    if any(p._p.xpath('.//w:br[@w:type="page"]') for p in doc.paragraphs[-1:]):
        raise ValueError('Prepared title must not end in a page break.')
    install_styles(doc)
    install_toc_styles(doc)
    doc.core_properties.title = 'Отчет по практике'
    doc.core_properties.author = ''; doc.core_properties.last_modified_by = ''
    doc.core_properties.subject = ''; doc.core_properties.comments = ''; doc.core_properties.keywords = ''
    for section in doc.sections:
        section.page_width = Mm(210); section.page_height = Mm(297)
        section.left_margin = Mm(30); section.right_margin = Mm(10)
        section.top_margin = section.bottom_margin = Mm(20)
    first = doc.sections[0]
    first.different_first_page_header_footer = True
    for el in list(first.first_page_footer._element):
        first.first_page_footer._element.remove(el)
    main = doc.add_section(WD_SECTION_START.NEW_PAGE)
    main.different_first_page_header_footer = False
    main.footer.is_linked_to_previous = False
    for el in list(main.footer._element):
        main.footer._element.remove(el)
    footer = main.footer.add_paragraph('', 'Practice Caption')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Mm(0)
    field(footer, 'PAGE', '')
    for section in doc.sections:
        for el in section._sectPr.findall(qn('w:pgNumType')):
            section._sectPr.remove(el)
    doc.add_paragraph('СОДЕРЖАНИЕ', 'Practice TOC Heading')
    toc = doc.add_paragraph('', 'Practice Body')
    toc.paragraph_format.first_line_indent = Mm(0)
    field(toc, 'TOC \\o "1-2" \\h \\z \\u', '')
    fields = OxmlElement('w:updateFields'); fields.set(qn('w:val'), 'true')
    doc.settings.element.append(fields)
    headings = {h['text']: h for h in plan['headings']}
    def blank(style, lines=1):
        paragraph = doc.add_paragraph('', style)
        # Set the paragraph mark as well: it controls empty-line metrics in Word.
        marker = OxmlElement('w:rPr')
        size = OxmlElement('w:sz'); size.set(qn('w:val'),'28'); marker.append(size)
        paragraph._p.get_or_add_pPr().append(marker)
        run = paragraph.add_run('\u00a0')
        run.font.name = 'Times New Roman'; run.font.size = Pt(14); run.font.hidden = False
        for _ in range(lines-1):
            run.add_break()
        return paragraph
    previous = None
    blocks = content.get('blocks', [])
    for index, block in enumerate(blocks):
        kind = block.get('type')
        if kind == 'heading':
            text = block['text']
            if text not in headings:
                raise ValueError('Heading absent from plan: ' + text)
            heading = headings[text]
            numbered = re.match(r'^\d+(?:\.\d+)*\s', text)
            if heading['level'] == 2:
                follows_section = (previous and previous.get('type') == 'heading'
                                   and headings[previous['text']]['level'] == 1)
                if not follows_section:
                    gap = blank('Practice Body')
                    # Move the required gap with its heading, without forcing a page.
                    gap.paragraph_format.keep_with_next = True
                style = 'Practice Subsection'
            else:
                style = 'Practice Section' if numbered else 'Practice Structural'
            doc.add_paragraph(text, style)
        elif kind == 'paragraph':
            doc.add_paragraph(normalize_dashes(block['text']), 'Practice Body')
        elif kind == 'page_break':
            # Legacy input is harmless only at a section boundary. The heading
            # style already supplies that break; never add a second one.
            following = blocks[index + 1] if index + 1 < len(blocks) else {}
            if (following.get('type') != 'heading'
                    or headings.get(following.get('text'), {}).get('level') != 1):
                raise ValueError('Explicit body page breaks are forbidden; use level-1 headings and automatic text flow.')
            continue
        elif kind == 'figure':
            number = block['number']
            if not re.fullmatch(r'\d+\.\d+', number):
                raise ValueError('Figure number must use section.item.')
            if block.get('status') == 'pending-manual':
                if block.get('file'):
                    raise ValueError('Pending figure cannot contain an image file.')
                lines = block.get('blank_lines', 5)
                if not isinstance(lines, int) or not 1 <= lines <= 12:
                    raise ValueError('blank_lines must be between 1 and 12.')
                blank('Practice Figure',lines)
            else:
                path = (root / block['file']).resolve()
                if not path.is_relative_to(root.resolve()) or not path.is_file():
                    raise ValueError('Figure file missing/outside workspace.')
                width = float(block.get('width_mm', 160))
                if not 0 < width <= 170:
                    raise ValueError('Figure width exceeds printable area.')
                p = doc.add_paragraph('', 'Practice Figure')
                p.add_run().add_picture(str(path), width=Mm(width))
            doc.add_paragraph('Рисунок ' + number + ' ' + normalize_dashes(block['title']), 'Practice Caption')
        elif kind == 'table':
            headers = block['headers']
            if not headers or any(len(row) != len(headers) for row in block['rows']):
                raise ValueError('Table rows must match headers.')
            doc.add_paragraph(normalize_dashes(block['caption']), 'Practice Table Caption')
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = block.get('widths_mm')
            if widths is not None:
                if len(widths) != len(headers) or any(not isinstance(w,(int,float)) or w<=0 for w in widths) or sum(widths)>170:
                    raise ValueError('Table widths must match columns and fit within 170 mm.')
                table.autofit = False
                for column, width in zip(table.columns,widths): column.width = Mm(width)
            repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cells, row in [(table.rows[0].cells, headers)]:
                for cell, text in zip(cells, row):
                    p = cell.paragraphs[0]; p.style = 'Practice Table'
                    p.add_run(normalize_dashes(str(text)))
            for row in block['rows']:
                for cell, text in zip(table.add_row().cells, row):
                    p = cell.paragraphs[0]; p.style = 'Practice Table'
                    p.add_run(normalize_dashes(str(text)))
            if widths is not None:
                for row in table.rows:
                    for cell,width in zip(row.cells,widths): cell.width = Mm(width)
        else:
            raise ValueError('Unsupported content block: ' + str(kind))
        previous = block
    issues = structure_issues(doc, plan)
    if issues:
        raise ValueError(issues[0]['message'])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ('title', 'content', 'plan', 'output'):
        parser.add_argument('--' + name, required=True, type=Path)
    args = parser.parse_args()
    data = json.loads(args.content.read_text(encoding='utf-8-sig'))
    print(build(args.title, data, load_plan(args.plan), args.output).resolve())
