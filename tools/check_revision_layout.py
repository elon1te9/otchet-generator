"""Formatting regression gate for report body; optional PDF and protected-prefix QA."""
import argparse, json, re
import hashlib
from pathlib import Path
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm
from report_pagination import effective, pagination_issues
from rendered_layout import inspect_rendered
from report_plan import (load_plan, body_start, body_headings, outline_level,
                         normalized, role_index, section_end, structure_issues, figure_prior, cached_toc)

def start_index(doc, plan=None):
    return body_start(doc, plan)

def font(p,r,attr):
    value=getattr(r.font,attr); s=r.style
    while value is None and s is not None:
        value=getattr(s.font,attr); s=s.base_style
    s=p.style
    while value is None and s is not None:
        value=getattr(s.font,attr); s=s.base_style
    return value

def protected(doc, plan=None):
    begin=start_index(doc,plan)
    if begin is None:
        raise ValueError('Cannot identify protected prefix boundary.')
    end=doc.paragraphs[begin]._p
    elements=[]
    for el in doc._element.body:
        if el is end: break
        # Ignore namespace layout differences; comparison is semantic XML.
        elements.append(etree.tostring(el,method='c14n',exclusive=True))
    return elements

def check(doc, pdf=None, baseline=None, plan=None, manifest=None):
    issues=[]; begin=start_index(doc,plan); pending=[]
    def issue(code,text): issues.append({'code':code,'message':text})
    if begin is None:
        return {'status':'fail','issues':[{'code':'body-start','message':'Body heading missing.'}],'pages':[]}
    if plan:
        issues.extend(structure_issues(doc,plan))
    issues.extend(pagination_issues(doc,begin))
    def para(p, label, size=14, spacing=1.5, indent=None, align=None):
        for attr in ('space_before','space_after','left_indent','right_indent'):
            if effective(p,attr) not in (None,0): issue('paragraph-spacing',f'{label}: {attr} must be zero')
        if abs(float(effective(p,'line_spacing') or 0)-spacing)>.01: issue('line-spacing',f'{label}: expected {spacing}')
        if indent is not None and abs((effective(p,'first_line_indent') or 0)-indent)>635: issue('indent',label)
        if align is not None and effective(p,'alignment')!=align: issue('alignment',label)
        for r in p.runs:
            if not r.text: continue
            z=font(p,r,'size')
            if z is None or abs(z.pt-size)>.1: issue('font-size',f'{label}: expected {size}')
            if font(p,r,'name')!='Times New Roman': issue('font-family',label)
    for i,p in enumerate(doc.paragraphs[begin:],begin):
        label=f'paragraph {i}: {p.text[:55]}'
        if re.search('[—–]',p.text): issue('long-dash',label)
        if re.search(r'\b(?:на рисунке|рисунка|рисунке)\s+\d',p.text,re.I): issue('figure-reference',label)
        style=p.style.name
        is_heading=outline_level(p) in (0,1)
        para(p,label)
        if is_heading:
            if any(font(p,r,'bold') for r in p.runs if r.text): issue('bold-heading',label)
            if any(font(p,r,'italic') or font(p,r,'underline') for r in p.runs if r.text): issue('heading-face',label)
            if outline_level(p)==0:
                if re.match(r'^\d+\s',p.text): para(p,label,indent=Mm(12.5),align=WD_ALIGN_PARAGRAPH.LEFT)
                else: para(p,label,indent=0,align=WD_ALIGN_PARAGRAPH.CENTER)
            if outline_level(p)==1:
                para(p,label,indent=Mm(12.5),align=WD_ALIGN_PARAGRAPH.LEFT)
        if not is_heading:
            is_gap=(not p.text.strip() and not p._p.xpath('.//w:drawing | .//w:br')
                    and i+1<len(doc.paragraphs) and outline_level(doc.paragraphs[i+1])==1)
            for attr in ('keep_together','keep_with_next','page_break_before'):
                if effective(p,attr) and not (is_gap and attr=='keep_with_next'):
                    issue('pagination-marker',f'{label}: {attr}')
            if p.text.strip() and not p.text.startswith(('Рисунок ', 'Таблица ')) and not p._p.xpath('.//w:drawing'):
                para(p,label,indent=Mm(12.5),align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        if p.text.startswith('Таблица '):
            para(p,label,indent=Mm(12.5),align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        if p.text.startswith('Рисунок ') or p._p.xpath('.//w:drawing'):
            para(p,label,indent=0)
        if p.text.startswith('Рисунок '):
            m=re.match(r'Рисунок (\d+\.\d+) ',p.text)
            prior=figure_prior(doc,i,begin)
            if not m or f'(Рисунок {m.group(1)})' not in prior: issue('figure-reference',label)
    body_elements=list(doc._element.body)
    body_anchor=body_elements.index(doc.paragraphs[begin]._p)
    body_tables=[t for t in doc.tables if body_elements.index(t._tbl)>=body_anchor]
    for ti,t in enumerate(body_tables):
        for ri,row in enumerate(t.rows):
            for ci,c in enumerate(row.cells):
                for p in c.paragraphs:
                    para(p,f'table {ti} row {ri} cell {ci}',12,1,0,WD_ALIGN_PARAGRAPH.JUSTIFY)
                    if any(effective(p,a) for a in ('keep_together','keep_with_next','page_break_before')):
                        issue('pagination-marker',f'table {ti} row {ri} cell {ci}')
                    if re.search('[—–]',p.text): issue('long-dash',f'table {ti}')
    if baseline:
        try:
            if protected(doc,plan)!=protected(Document(baseline),plan): issue('protected-prefix','Title/TOC XML differs from baseline')
        except ValueError as error:
            issue('protected-prefix',str(error))
    guide=role_index(doc,begin,'guide',plan)
    testing=role_index(doc,begin,'testing',plan)
    captions={p.text.strip():i for i,p in enumerate(doc.paragraphs) if i>=begin and p.text.startswith('Рисунок ')}
    pending_captions=set()
    if plan and captions and manifest is None:
        issue('manifest-required','Figures require the current evidence manifest.')
    if manifest is not None:
        data=json.loads(Path(manifest).read_text(encoding='utf-8-sig')) if not isinstance(manifest,dict) else manifest
        if not plan or not plan.get('input_signature') or data.get('input_signature')!=plan['input_signature']:
            issue('manifest-inputs','Manifest must match the current plan input signature.')
        records=data.get('figures',[])
        if not isinstance(records,list):
            issue('manifest-format','figures must be a list'); records=[]
        recorded=set()
        for record in records:
            if not isinstance(record,dict):
                issue('manifest-format','Figure entry must be a mapping'); continue
            caption=record.get('caption','')
            if caption in recorded: issue('manifest-duplicate',caption)
            recorded.add(caption)
            match=re.match(r'^Рисунок (\d+\.\d+) ',caption)
            if not match or match[1]!=record.get('number') or caption not in captions:
                issue('manifest-caption','Missing or inconsistent caption: '+caption); continue
            index=captions[caption]
            current_heading=next((p.text for i,p in reversed(body_headings(doc,begin)) if i<index),None)
            if record.get('section')!=current_heading:
                issue('manifest-section',caption)
            if not all(isinstance(record.get(k),str) and record[k].strip() for k in ('scenario','source','log')):
                issue('manifest-evidence',caption)
            if record.get('status')=='pending-manual':
                if record.get('file') or not all(record.get(k) for k in ('expected_state','blocker')):
                    issue('pending-evidence',caption); continue
                preceding=doc.paragraphs[index-1] if index>begin else None
                if preceding is None or preceding.text.strip() or preceding._p.xpath('.//w:drawing'):
                    issue('pending-area',caption); continue
                pending.append(caption); pending_captions.add(caption)
            elif record.get('status')=='observed':
                review=record.get('capture_review',{})
                if (not isinstance(review,dict) or review.get('cursor') not in ('neutral','intentional-hover')
                        or review.get('ready') is not True or review.get('readable') is not True
                        or not isinstance(review.get('note'),str) or not review['note'].strip()):
                    issue('capture-review',caption + ': inspect the original screenshot and record cursor/readiness/readability.')
                preceding=doc.paragraphs[index-1] if index>begin else None
                if not record.get('file') or preceding is None or not preceding._p.xpath('.//w:drawing'):
                    issue('observed-image',caption)
                else:
                    root=Path(__file__).resolve().parents[1]
                    image_path=(root/record['file']).resolve()
                    if not image_path.is_relative_to(root) or not image_path.is_file():
                        issue('observed-image-source',caption)
                    else:
                        links=preceding._p.xpath('.//a:blip/@r:embed')
                        expected_hash=hashlib.sha256(image_path.read_bytes()).hexdigest()
                        actual_hashes=[hashlib.sha256(doc.part.related_parts[link].blob).hexdigest() for link in links]
                        if expected_hash not in actual_hashes: issue('observed-image-source',caption)
            else:
                issue('manifest-status',caption)
        for caption in captions.keys()-recorded:
            issue('manifest-missing',caption)
    if guide is not None:
        end=section_end(doc,guide)
        has_image=any(p._p.xpath('.//w:drawing') for p in doc.paragraphs[guide:end])
        has_pending=any(guide<i<end and caption in pending_captions for caption,i in captions.items())
        if not has_image and not has_pending:
            issue('guide-figures','User guide has no observed images or documented pending areas')
    if testing is not None:
        testing_el=doc.paragraphs[testing]._p
        end=section_end(doc,testing)
        end_el=doc.paragraphs[end]._p if end<len(doc.paragraphs) else None
        following=list(doc._element.body)[list(doc._element.body).index(testing_el):]
        tables=[]
        for el in following:
            if el is end_el: break
            if el.tag==qn('w:tbl'): tables.append(el)
        expected=['№','Что проверялось','Ожидаемый результат','Результат']
        if len(tables)<2: issue('testing-tables','Separate black-box and white-box tables required')
        for t in tables:
            first=t.xpath('./w:tr')[0]
            actual=[''.join(c.xpath('.//w:t/text()')) for c in first.xpath('./w:tc')]
            if actual!=expected: issue('testing-columns',str(actual))
    pages=[]
    if pdf:
        rendered_headings=plan['headings'] if plan else [
            {'text':p.text,'level':outline_level(p)+1} for _,p in body_headings(doc,begin)]
        pages,render_issues,first_body=inspect_rendered(pdf,rendered_headings)
        issues.extend(render_issues)
        import pdfplumber
        with pdfplumber.open(pdf) as rendered:
            if len(rendered.pages)>25: issue('page-limit',str(len(rendered.pages)))
            texts=[page.extract_text() or '' for page in rendered.pages]
            def heading_pages(title):
                pattern=r'(?m)^' + r'\s+'.join(re.escape(word) for word in title.split()) + r'\s*$'
                return [i for i,text in enumerate(texts,1) if re.search(pattern,text)]
            # Word can recalculate PAGEREF on read-only PDF export. Compare with
            # the protected cached DOCX entries, not that recalculated PDF text.
            toc_entries=cached_toc(doc,begin)
            rendered_toc=normalized(' '.join(texts[:first_body-1])) if first_body else ''
            for _,p in body_headings(doc,begin):
                # Normalized PDF text permits line-wrapped headings.
                expected=toc_entries.get(normalized(p.text))
                actual=[i for i in heading_pages(p.text) if first_body and i>=first_body]
                if not expected: issue('toc-entry',f'Heading absent in cached TOC: {p.text}')
                elif not actual or actual[0]!=expected: issue('toc-page',f'{p.text}: expected {expected}, actual {actual}')
                shown=re.search(re.escape(' '.join(p.text.split()))+r'[\s.]+(\d+)\b',rendered_toc)
                if expected and (not shown or int(shown.group(1))!=expected):
                    issue('toc-render-page',f'{p.text}: cached {expected}, rendered {shown.group(1) if shown else "missing"}; inspect body bookmark target')
    return {'status':'fail' if issues else 'pending-manual' if pending else 'pass',
            'issues':issues,'pending_manual':pending,'pages':pages,
            'rendered_status':('fail' if any(i['code'].startswith(('pdf-','page-','toc-')) for i in issues) else 'pass') if pdf else 'not-run',
            'complete_verification':False,
            'visual_review_required':True,'evidence_review_required':True}

if __name__=='__main__':
    from report_cli import add_summary_argument, require_summary_artifact, emit_result
    a=argparse.ArgumentParser(); a.add_argument('docx',type=Path); a.add_argument('--plan',required=True,type=Path); a.add_argument('--manifest',type=Path); a.add_argument('--pdf',type=Path); a.add_argument('--baseline',type=Path); a.add_argument('--json',type=Path)
    add_summary_argument(a); args=a.parse_args(); require_summary_artifact(a,args)
    result=check(Document(args.docx),args.pdf,args.baseline,load_plan(args.plan),args.manifest)
    emit_result(result,args.json,args.summary)
    raise SystemExit(1 if result['status']=='fail' else 2 if result['status']=='pending-manual' else 0)
