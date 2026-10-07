#!/usr/bin/env python3
"""Read-only DOCX inspector for report templates and completed reports."""

from __future__ import annotations

import argparse
import json
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


def emu_to_mm(value):
    return None if value is None else round(value / 36000, 2)


def pt(value):
    return None if value is None else round(value.pt, 2)


def alignment_name(value):
    return None if value is None else str(value)


def paragraph_record(p, index, location="body"):
    fmt = p.paragraph_format
    runs = []
    for run in p.runs:
        if not run.text and not run._r.xpath(".//w:drawing|.//w:pict|.//w:fldChar|.//w:instrText"):
            continue
        runs.append(
            {
                "text": run.text,
                "font": run.font.name,
                "size_pt": pt(run.font.size),
                "bold": run.bold,
                "italic": run.italic,
                "underline": bool(run.underline) if run.underline is not None else None,
            }
        )
    xml = p._p
    return {
        "index": index,
        "location": location,
        "text": p.text,
        "style": p.style.name if p.style else None,
        "alignment": alignment_name(p.alignment),
        "left_indent_mm": emu_to_mm(fmt.left_indent),
        "right_indent_mm": emu_to_mm(fmt.right_indent),
        "first_line_indent_mm": emu_to_mm(fmt.first_line_indent),
        "space_before_pt": pt(fmt.space_before),
        "space_after_pt": pt(fmt.space_after),
        "line_spacing": (
            round(fmt.line_spacing, 3)
            if isinstance(fmt.line_spacing, float)
            else pt(fmt.line_spacing)
        ),
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
        "has_page_break": bool(xml.xpath(".//w:br[@w:type='page']")),
        "has_section_properties": bool(xml.xpath("./w:pPr/w:sectPr")),
        "field_codes": [n.text or "" for n in xml.xpath(".//w:instrText")],
        "has_drawing": bool(xml.xpath(".//w:drawing|.//w:pict")),
        "runs": runs,
    }


def style_record(style):
    pf = getattr(style, "paragraph_format", None)
    font = getattr(style, "font", None)
    base_style = getattr(style, "base_style", None)
    return {
        "name": style.name,
        "type": str(style.type),
        "base_style": base_style.name if base_style else None,
        "font": None if font is None else {
            "name": font.name,
            "size_pt": pt(font.size),
            "bold": font.bold,
            "italic": font.italic,
        },
        "paragraph": None if pf is None else {
            "alignment": alignment_name(pf.alignment),
            "left_indent_mm": emu_to_mm(pf.left_indent),
            "right_indent_mm": emu_to_mm(pf.right_indent),
            "first_line_indent_mm": emu_to_mm(pf.first_line_indent),
            "space_before_pt": pt(pf.space_before),
            "space_after_pt": pt(pf.space_after),
            "line_spacing": round(pf.line_spacing, 3) if isinstance(pf.line_spacing, float) else pt(pf.line_spacing),
            "keep_with_next": pf.keep_with_next,
            "page_break_before": pf.page_break_before,
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--json", type=Path)
    parser.add_argument("--text", type=Path)
    args = parser.parse_args()

    doc = Document(args.input)
    paragraphs = [paragraph_record(p, i) for i, p in enumerate(doc.paragraphs)]
    table_rows = []
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            table_rows.append({
                "table": ti,
                "row": ri,
                "cells": [cell.text for cell in row.cells],
                "repeat_header": bool(row._tr.xpath("./w:trPr/w:tblHeader")),
            })

    headers_footers = []
    sections = []
    for i, sec in enumerate(doc.sections):
        sections.append({
            "index": i,
            "start_type": str(sec.start_type),
            "page_width_mm": emu_to_mm(sec.page_width),
            "page_height_mm": emu_to_mm(sec.page_height),
            "orientation": str(sec.orientation),
            "top_margin_mm": emu_to_mm(sec.top_margin),
            "bottom_margin_mm": emu_to_mm(sec.bottom_margin),
            "left_margin_mm": emu_to_mm(sec.left_margin),
            "right_margin_mm": emu_to_mm(sec.right_margin),
            "header_distance_mm": emu_to_mm(sec.header_distance),
            "footer_distance_mm": emu_to_mm(sec.footer_distance),
            "different_first_page": sec.different_first_page_header_footer,
            "header_linked": sec.header.is_linked_to_previous,
            "footer_linked": sec.footer.is_linked_to_previous,
        })
        for kind, part in (("header", sec.header), ("footer", sec.footer),
                           ("first_page_header", sec.first_page_header),
                           ("first_page_footer", sec.first_page_footer)):
            for pi, p in enumerate(part.paragraphs):
                rec = paragraph_record(p, pi, f"section_{i}_{kind}")
                if rec["text"] or rec["field_codes"] or rec["has_drawing"]:
                    headers_footers.append(rec)

    with zipfile.ZipFile(args.input) as zf:
        names = zf.namelist()
        settings = zf.read("word/settings.xml").decode("utf-8", errors="replace") if "word/settings.xml" in names else ""
        doc_xml_bytes = zf.read("word/document.xml")
        doc_xml = doc_xml_bytes.decode("utf-8", errors="replace")
        root = etree.fromstring(doc_xml_bytes)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        xml_paragraphs = []
        for index, element in enumerate(root.xpath(".//w:body//w:p", namespaces=ns)):
            text_value = "".join(element.xpath(".//w:t/text()", namespaces=ns))
            field_values = [v.strip() for v in element.xpath(".//w:instrText/text()", namespaces=ns)]
            style_values = element.xpath("./w:pPr/w:pStyle/@w:val", namespaces=ns)
            if text_value.strip() or field_values:
                xml_paragraphs.append({
                    "index": index,
                    "text": text_value,
                    "field_codes": field_values,
                    "style_id": style_values[0] if style_values else None,
                    "inside_content_control": bool(element.xpath("ancestor::w:sdt", namespaces=ns)),
                })
        package = {
            "parts": len(names),
            "media": [n for n in names if n.startswith("word/media/")],
            "headers": [n for n in names if n.startswith("word/header") and n.endswith(".xml")],
            "footers": [n for n in names if n.startswith("word/footer") and n.endswith(".xml")],
            "has_comments": "word/comments.xml" in names,
            "has_numbering": "word/numbering.xml" in names,
            "track_revisions": "<w:trackRevisions" in settings,
            "update_fields_on_open": "<w:updateFields" in settings,
            "field_codes": Counter(
                code.strip().split()[0] if code.strip() else ""
                for code in [
                    *sum((p["field_codes"] for p in xml_paragraphs), []),
                    *sum((p["field_codes"] for p in headers_footers), []),
                ]
            ),
            "inline_images": doc_xml.count("<wp:inline"),
            "floating_images": doc_xml.count("<wp:anchor"),
        }

    result = {
        "file": str(args.input.resolve()),
        "core_properties": {
            "title": doc.core_properties.title,
            "subject": doc.core_properties.subject,
            "author": doc.core_properties.author,
            "last_modified_by": doc.core_properties.last_modified_by,
            "created": str(doc.core_properties.created) if doc.core_properties.created else None,
            "modified": str(doc.core_properties.modified) if doc.core_properties.modified else None,
        },
        "counts": {
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "sections": len(doc.sections),
            "images": len(package["media"]),
        },
        "sections": sections,
        "styles": [style_record(s) for s in doc.styles],
        "paragraphs": paragraphs,
        "xml_paragraphs": xml_paragraphs,
        "table_rows": table_rows,
        "headers_footers": headers_footers,
        "package": package,
    }

    payload = json.dumps(result, ensure_ascii=False, indent=2, default=dict)
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(payload, encoding="utf-8")
    else:
        print(payload)

    if args.text:
        lines = []
        for p in paragraphs:
            if p["text"].strip():
                lines.append(f"[P{p['index']}|{p['style']}] {p['text']}")
        for row in table_rows:
            lines.append(f"[T{row['table']}R{row['row']}] " + " | ".join(row["cells"]))
        for p in headers_footers:
            if p["text"].strip() or p["field_codes"]:
                lines.append(f"[{p['location']}] {p['text']} {' '.join(p['field_codes'])}".rstrip())
        args.text.parent.mkdir(parents=True, exist_ok=True)
        args.text.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
