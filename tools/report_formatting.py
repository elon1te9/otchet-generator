"""Report-only styles; never change shared title/TOC styles when revising a DOCX."""
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def install_styles(doc):
    definitions = {
        'Practice Body': (14, 1.5, 12.5, WD_ALIGN_PARAGRAPH.JUSTIFY, None),
        'Practice Section': (14, 1.5, 12.5, WD_ALIGN_PARAGRAPH.LEFT, 0),
        'Practice Subsection': (14, 1.5, 12.5, WD_ALIGN_PARAGRAPH.LEFT, 1),
        'Practice Structural': (14, 1.5, 0, WD_ALIGN_PARAGRAPH.CENTER, 0),
        'Practice TOC Heading': (14, 1.5, 0, WD_ALIGN_PARAGRAPH.CENTER, None),
        'Practice Figure': (14, 1.5, 0, WD_ALIGN_PARAGRAPH.CENTER, None),
        'Practice Caption': (14, 1.5, 0, WD_ALIGN_PARAGRAPH.CENTER, None),
        'Practice Table Caption': (14, 1.5, 12.5, WD_ALIGN_PARAGRAPH.JUSTIFY, None),
        'Practice Table': (12, 1, 0, WD_ALIGN_PARAGRAPH.JUSTIFY, None),
    }
    for name, (size, spacing, indent, align, outline) in definitions.items():
        s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = doc.styles['Normal']
        s.font.name = 'Times New Roman'; s.font.size = Pt(size)
        s.font.bold = False; s.font.italic = False; s.font.underline = False
        s.font.color.rgb = RGBColor(0, 0, 0)
        fonts = s.element.get_or_add_rPr().get_or_add_rFonts()
        for k in list(fonts.attrib): del fonts.attrib[k]
        for k in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+k), 'Times New Roman')
        pf = s.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.first_line_indent = Mm(indent); pf.left_indent = pf.right_indent = Mm(0)
        pf.line_spacing = spacing; pf.alignment = align
        # Pagination flags create visible margin squares in Word's Show All view.
        # Ordinary paragraphs use explicit, sentence-aligned page breaks instead.
        pf.keep_together = outline is not None; pf.widow_control = True
        pf.keep_with_next = outline is not None
        pf.page_break_before = False
        if outline is not None:
            for old in s.element.xpath('./w:pPr/w:outlineLvl'):
                old.getparent().remove(old)
            e = OxmlElement('w:outlineLvl'); e.set(qn('w:val'), str(outline)); s.element.get_or_add_pPr().append(e)

def normalize_dashes(text):
    return text.replace('\u2014', '-').replace('\u2013', '-')

def install_toc_styles(doc):
    """New documents only; never apply to a protected user's cached contents."""
    for level in (1, 2):
        name = f'toc {level}'
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
        # Word renames custom styles that collide with built-in TOC names.
        style.element.attrib.pop(qn('w:customStyle'), None)
        style.base_style = doc.styles['Normal']
        style.font.name = 'Times New Roman'; style.font.size = Pt(14)
        style.font.bold = False; style.font.italic = False; style.font.underline = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        pf = style.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.5; pf.first_line_indent = Mm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Mm(0 if level == 1 else 12.5); pf.right_indent = Mm(0)
        pf.keep_together = False; pf.keep_with_next = False; pf.page_break_before = False
        for node in style.element.xpath('./w:pPr/w:outlineLvl'):
            node.getparent().remove(node)

def field(paragraph, code, cached):
    for kind in ('begin', 'separate', 'end'):
        if kind == 'separate':
            r = paragraph.add_run(); e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve'); e.text = ' '+code+' '; r._r.append(e)
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
        paragraph.add_run()._r.append(e)
        if kind == 'separate': paragraph.add_run(cached)
