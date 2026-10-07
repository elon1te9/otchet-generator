"""Shared report structure; no project-specific heading lists."""
import json
import re
from pathlib import Path
from docx.oxml.ns import qn

ROLES = {'assignment', 'introduction', 'section', 'subsection', 'guide',
         'testing', 'conclusion', 'appendix'}

def normalized(text):
    return ' '.join(text.split())

def is_toc(paragraph):
    return paragraph.style.name.lower().startswith(('toc', 'содержание', 'оглавление'))

def outline_level(paragraph):
    nodes = paragraph._p.xpath('./w:pPr/w:outlineLvl')
    style = paragraph.style
    while not nodes and style is not None:
        nodes = style.element.xpath('./w:pPr/w:outlineLvl')
        style = style.base_style
    if nodes:
        return int(nodes[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'))
    match = re.search(r'(?:Heading|Заголовок)\s+(\d+)', paragraph.style.name, re.I)
    return int(match[1]) - 1 if match else None

def load_plan(path):
    if path is None:
        return None
    data = json.loads(Path(path).read_text(encoding='utf-8-sig'))
    if not isinstance(data, dict) or not isinstance(data.get('headings'), list) or not data['headings']:
        raise ValueError('Plan must contain a nonempty headings list.')
    seen = set()
    for heading in data['headings']:
        if not isinstance(heading, dict) or not isinstance(heading.get('text'), str) or not heading['text'].strip():
            raise ValueError('Every heading needs text.')
        key = normalized(heading['text'])
        if key in seen:
            raise ValueError('Duplicate heading: ' + key)
        seen.add(key)
        if heading.get('level') not in (1, 2) or heading.get('role') not in ROLES:
            raise ValueError('Heading needs level 1/2 and a supported role: ' + key)
    if data.get('body_start') != data['headings'][0]['text']:
        raise ValueError('body_start must equal the first heading.')
    for role in ('assignment', 'testing', 'conclusion'):
        if sum(h['role'] == role for h in data['headings']) != 1:
            raise ValueError('Plan must contain exactly one ' + role + ' heading.')
    return data

def body_start(doc, plan=None):
    if plan:
        return next((i for i, p in enumerate(doc.paragraphs)
                     if not is_toc(p) and normalized(p.text) == normalized(plan['body_start'])), None)
    # Diagnostics without a plan: use an actual body heading, never a cached TOC entry.
    return next((i for i, p in enumerate(doc.paragraphs)
                 if not is_toc(p) and p.text.strip() != 'СОДЕРЖАНИЕ'
                 and outline_level(p) in (0, 1)), None)

def body_headings(doc, begin):
    return [(i, p) for i, p in enumerate(doc.paragraphs) if i >= begin
            and not is_toc(p) and outline_level(p) in (0, 1)]

def role_index(doc, begin, role, plan=None):
    titles = [h['text'] for h in plan['headings'] if h['role'] == role] if plan else []
    for i, p in body_headings(doc, begin):
        if titles and normalized(p.text) in map(normalized, titles):
            return i
        if not plan:
            title = re.sub(r'^\d+(?:\.\d+)*\s+', '', p.text).strip().casefold()
            if role == 'guide' and 'руководство пользователя' in title:
                return i
            if role == 'testing' and title == 'тестирование':
                return i
            if role == 'conclusion' and title == 'заключение':
                return i
    return None

def section_end(doc, index):
    level = outline_level(doc.paragraphs[index])
    return next((i for i, p in body_headings(doc, index + 1)
                 if outline_level(p) <= level), len(doc.paragraphs))

def structure_issues(doc, plan):
    begin = body_start(doc, plan)
    if begin is None:
        return [{'code': 'body-start', 'message': 'First body heading missing.'}]
    actual = [(normalized(p.text), outline_level(p) + 1) for _, p in body_headings(doc, begin)]
    expected = [(normalized(h['text']), h['level']) for h in plan['headings']]
    return [] if actual == expected else [{'code': 'heading-plan', 'message':
        f'Body headings differ from plan. Expected {expected}; actual {actual}'}]

def figure_prior(doc, index, begin):
    texts = []
    for paragraph in reversed(doc.paragraphs[begin:index]):
        if outline_level(paragraph) in (0, 1):
            break
        if paragraph.text.strip():
            texts.append(paragraph.text)
        if len(texts) == 4:
            break
    return ' '.join(reversed(texts))

def cached_toc(doc, begin):
    """Read cached entries, including Word's TOC content control (w:sdt)."""
    entries = {}
    style_names = {style.style_id: style.name for style in doc.styles}
    anchor = doc.paragraphs[begin]._p
    for element in doc._element.body:
        if element is anchor:
            break
        paragraphs = [element] if element.tag == qn('w:p') else list(element.iter(qn('w:p')))
        for paragraph in paragraphs:
            style = paragraph.find(qn('w:pPr') + '/' + qn('w:pStyle'))
            style_id = style.get(qn('w:val'), '') if style is not None else ''
            if not style_names.get(style_id, style_id).casefold().startswith(('toc', 'содержание', 'оглавление')):
                continue
            chunks = []
            for node in paragraph.iter():
                if node.tag.endswith('}t'):
                    chunks.append(node.text or '')
                elif node.tag.endswith('}tab'):
                    chunks.append('\t')
            match = re.match(r'^(.*?)\t(\d+)\s*$', ''.join(chunks))
            if match:
                entries[normalized(match[1])] = int(match[2])
    return entries
