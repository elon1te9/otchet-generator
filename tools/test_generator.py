"""Generator integration checks using isolated synthetic inputs, never project claims."""
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import yaml

from build_report import build
from check_revision_layout import check as check_layout
from check_subsection_conclusions import check as check_conclusions
from package_generator import package
from prepare_report import prepare
from report_formatting import install_styles
from report_plan import load_plan, cached_toc, body_start, structure_issues

ROOT = Path(__file__).resolve().parents[1]

def synthetic_plan():
    return {'body_start': 'ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ НА ТЕМУ: Проверка генератора',
            'input_signature': 'synthetic-fixture-only', 'headings': [
        {'text': 'ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ НА ТЕМУ: Проверка генератора', 'level': 1, 'role': 'assignment'},
        {'text': '1 УСТРОЙСТВО ПРОГРАММЫ', 'level': 1, 'role': 'section'},
        {'text': '1.1 Обработка файлов', 'level': 2, 'role': 'subsection'},
        {'text': '1.2 Порядок работы', 'level': 2, 'role': 'guide'},
        {'text': '2 ТЕСТИРОВАНИЕ', 'level': 1, 'role': 'testing'},
        {'text': 'ЗАКЛЮЧЕНИЕ', 'level': 1, 'role': 'conclusion'},
    ]}

def synthetic_content():
    plan=synthetic_plan(); h=plan['headings']
    return {'blocks': [
        {'type': 'heading', 'text': h[0]['text']},
        {'type': 'paragraph', 'text': 'Этот документ является тестовым входом для проверки генератора. Он не описывает реальное приложение и не подтверждает испытания пользовательского проекта.'},
        {'type': 'heading', 'text': h[1]['text']},
        {'type': 'heading', 'text': h[2]['text']},
        {'type': 'paragraph', 'text': 'Для проверки сменной структуры выбран подраздел с новым названием. Структурная проверка должна сопоставить его с планом.'},
        {'type': 'paragraph', 'text': 'Таким образом, тестовый подраздел задает условие проверки порядка заголовков генератора.'},
        {'type': 'heading', 'text': h[3]['text']},
        {'type': 'paragraph', 'text': 'Для проверки запасного оформления предусмотрена пустая область (Рисунок 1.1).'},
        {'type': 'figure', 'number': '1.1', 'title': 'Пустая область тестового документа', 'status': 'pending-manual', 'blank_lines': 3},
        {'type': 'paragraph', 'text': 'В результате, тестовый документ содержит место для ручной вставки изображения без имитации интерфейса.'},
        {'type': 'page_break'},
        {'type': 'heading', 'text': h[4]['text']},
        {'type': 'paragraph', 'text': 'Ниже размещены синтетические таблицы для проверки оформления черного и белого ящика. Содержимое таблиц не является результатами испытаний приложения.'},
        {'type': 'table', 'caption': 'Таблица 2.1 - Черный ящик', 'headers': ['№', 'Что проверялось', 'Ожидаемый результат', 'Результат'],
         'rows': [['1', 'Тестовая строка', 'Четыре графы', 'Синтетический пример']]},
        {'type': 'paragraph', 'text': 'Таблица используется исключительно для проверки стиля ячеек генератора.'},
        {'type': 'table', 'caption': 'Таблица 2.2 - Белый ящик', 'headers': ['№', 'Что проверялось', 'Ожидаемый результат', 'Результат'],
         'rows': [['1', 'Тестовая строка', 'Отдельная таблица', 'Синтетический пример']]},
        {'type': 'paragraph', 'text': 'В итоге, документ предоставляет отдельные таблицы с установленными графами.'},
        {'type': 'heading', 'text': h[5]['text']},
        {'type': 'paragraph', 'text': 'Этот документ используется для проверки сборки и рендера генератора. Результаты проверки фиксируются отдельно в журналах, а не приписываются пользовательскому приложению.'},
    ]}

def synthetic_manifest():
    return {'input_signature': 'synthetic-fixture-only', 'figures': [{
        'number': '1.1', 'caption': 'Рисунок 1.1 Пустая область тестового документа',
        'section': '1.2 Порядок работы', 'status': 'pending-manual',
        'scenario': 'Синтетический вход проверки fallback', 'source': 'tools/test_generator.py',
        'log': 'temp/qa/tests.log', 'blocker': 'Иллюстрация намеренно не задана в тестовом входе',
        'expected_state': 'Пустая область тестового документа',
    }]}

class GeneratorTests(unittest.TestCase):
    def setUp(self):
        (ROOT/'temp').mkdir(exist_ok=True)
        self.workspace=tempfile.TemporaryDirectory(dir=ROOT/'temp')
        self.root=Path(self.workspace.name)
        self.plan=synthetic_plan()
        self.title=self.root/'title.docx'
        title=Document(); title.add_paragraph('Тестовый титул генератора', 'Title'); title.save(self.title)
        self.report=self.root/'report.docx'
        build(self.title,synthetic_content(),self.plan,self.report,self.root)
        self.doc=Document(self.report)
    def tearDown(self):
        self.workspace.cleanup()
    def codes(self,manifest=None):
        return {item['code'] for item in check_layout(self.doc,plan=self.plan,manifest=manifest)['issues']}
    def test_new_structure_and_subsection_conclusions(self):
        self.assertEqual(structure_issues(self.doc,self.plan),[])
        self.assertEqual(check_conclusions(self.doc,self.plan)['status'],'pass')
        self.assertEqual(self.codes(synthetic_manifest()),set())
    def test_changed_order_is_rejected(self):
        self.plan['headings'][2],self.plan['headings'][3]=self.plan['headings'][3],self.plan['headings'][2]
        self.assertIn('heading-plan',self.codes(synthetic_manifest()))
    def test_missing_subsection_is_rejected(self):
        for p in self.doc.paragraphs:
            if p.text=='1.1 Обработка файлов':
                p._p.getparent().remove(p._p); break
        self.assertEqual(check_conclusions(self.doc,self.plan)['status'],'fail')
    def test_pending_area_is_not_a_completed_report(self):
        result=check_layout(self.doc,plan=self.plan,manifest=synthetic_manifest())
        self.assertEqual(result['status'],'pending-manual')
        self.assertEqual(len(result['pending_manual']),1)
    def test_blank_area_requires_specific_manifest(self):
        self.assertIn('manifest-required',self.codes())
        self.assertIn('guide-figures',self.codes())
        manifest=synthetic_manifest(); manifest['figures'][0]['caption']='Рисунок 9.9 Чужой рисунок'
        self.assertIn('manifest-caption',self.codes(manifest))
        self.assertIn('guide-figures',self.codes(manifest))
    def test_stale_evidence_is_rejected(self):
        manifest=synthetic_manifest(); manifest['input_signature']='another-project'
        self.assertIn('manifest-inputs',self.codes(manifest))
    def test_observed_status_requires_image(self):
        manifest=synthetic_manifest(); manifest['figures'][0]['status']='observed'; manifest['figures'][0]['file']='missing.png'
        self.assertIn('observed-image',self.codes(manifest))
    def test_screenshot_review_requires_actual_quality_checks(self):
        manifest=synthetic_manifest(); record=manifest['figures'][0]
        record.update(status='observed',file='missing.png')
        self.assertIn('capture-review',self.codes(manifest))
        record['capture_review']={'cursor':'neutral','ready':True,'readable':True,'note':'Synthetic manifest validation input.'}
        self.assertNotIn('capture-review',self.codes(manifest))
        # A checklist cannot substitute for the actual image.
        self.assertIn('observed-image',self.codes(manifest))
        record['capture_review']['ready']=False
        self.assertIn('capture-review',self.codes(manifest))
    def test_toc_entry_inside_word_content_control(self):
        control=OxmlElement('w:sdt'); contents=OxmlElement('w:sdtContent'); control.append(contents)
        p=OxmlElement('w:p'); props=OxmlElement('w:pPr'); style=OxmlElement('w:pStyle'); style.set(qn('w:val'),'TOC1'); props.append(style); p.append(props)
        for text in ['1 УСТРОЙСТВО ПРОГРАММЫ',None,'3']:
            run=OxmlElement('w:r'); node=OxmlElement('w:tab' if text is None else 'w:t'); node.text=text; run.append(node); p.append(run)
        contents.append(p)
        anchor=self.doc.paragraphs[body_start(self.doc,self.plan)]._p; anchor.addprevious(control)
        self.assertEqual(cached_toc(self.doc,body_start(self.doc,self.plan))['1 УСТРОЙСТВО ПРОГРАММЫ'],3)
    def test_protected_prefix_change_is_rejected(self):
        self.doc.paragraphs[0].add_run(' Изменение.')
        result=check_layout(self.doc,baseline=self.report,plan=self.plan,manifest=synthetic_manifest())
        self.assertIn('protected-prefix',{i['code'] for i in result['issues']})
    def test_plan_requires_testing_and_unique_headings(self):
        path=self.root/'plan.json'; plan=synthetic_plan(); plan['headings'].pop(4)
        path.write_text(json.dumps(plan),encoding='utf-8')
        with self.assertRaises(ValueError): load_plan(path)
        plan=synthetic_plan(); plan['headings'].append(plan['headings'][1])
        path.write_text(json.dumps(plan),encoding='utf-8')
        with self.assertRaises(ValueError): load_plan(path)
    def test_toc_styles_are_built_in_for_word(self):
        self.assertFalse(self.doc.styles['toc 1'].element.get(qn('w:customStyle')))
    def test_new_document_uses_only_heading_page_breaks(self):
        from report_plan import outline_level
        from report_pagination import effective
        for p in self.doc.paragraphs:
            self.assertFalse(p._p.xpath('.//w:br[@w:type="page"]'))
            if outline_level(p) == 0:
                self.assertTrue(effective(p,'page_break_before'))
            if outline_level(p) == 1:
                self.assertFalse(effective(p,'page_break_before'))
    def test_explicit_break_before_subsection_or_text_is_rejected(self):
        content=synthetic_content()
        for target in (3,4):
            bad={'blocks':content['blocks'][:target]+[{'type':'page_break'}]+content['blocks'][target:]}
            with self.assertRaisesRegex(ValueError,'Explicit body page breaks'):
                build(self.title,bad,self.plan,self.report,self.root)
    def test_subsection_gap_after_unnumbered_level_one(self):
        plan=synthetic_plan(); plan['headings'][1]['text']='УСТРОЙСТВО ПРОГРАММЫ'
        content=synthetic_content(); content['blocks'][2]['text']=plan['headings'][1]['text']
        build(self.title,content,plan,self.report,self.root)
        doc=Document(self.report)
        index=next(i for i,p in enumerate(doc.paragraphs) if p.text=='1.1 Обработка файлов')
        self.assertEqual(doc.paragraphs[index-1].text,'УСТРОЙСТВО ПРОГРАММЫ')
    def test_word_localized_toc_ids(self):
        style=self.doc.styles['toc 1']; style.element.set(qn('w:styleId'),'14')
        paragraph=self.doc.paragraphs[body_start(self.doc,self.plan)].insert_paragraph_before('1 УСТРОЙСТВО ПРОГРАММЫ\t3')
        paragraph.style=style
        self.assertEqual(cached_toc(self.doc,body_start(self.doc,self.plan))['1 УСТРОЙСТВО ПРОГРАММЫ'],3)
    def test_packaging_excludes_all_personal_inputs(self):
        output=self.root/'generator.zip'; package(ROOT,output)
        with zipfile.ZipFile(output) as archive:
            self.assertIsNone(archive.testzip())
            names=archive.namelist()
            self.assertFalse(any(name.startswith(('temp/','output/','project/','screenshots/','references/current-example/','references/previous-reports/')) and not name.endswith('/.gitkeep') for name in names))
            config=yaml.safe_load(archive.read('config/student.yaml'))
            self.assertEqual(config['student']['full_name'],'')
            self.assertTrue(any(name.startswith('references/official/') and name.endswith('.pdf') for name in names))
    def test_config_dates_and_project_change(self):
        for directory in ('config','docs','skills','tools','references/current-example','references/official','project/example'):
            (self.root/directory).mkdir(parents=True,exist_ok=True)
        (self.root/'references/current-example/sample.txt').write_text('synthetic example',encoding='utf-8')
        (self.root/'references/official/rules.txt').write_text('synthetic rules',encoding='utf-8')
        source=self.root/'project/example/main.py'; source.write_text('value = 1',encoding='utf-8')
        assignment=self.root/'project/example/assignment.md'; assignment.write_text('synthetic assignment',encoding='utf-8')
        config={group:{field:'synthetic' for field in fields} for group,fields in {
            'student':['full_name','group','specialty_code','specialty_name'], 'institution':['name','city'],
            'practice':['kind','module_code','practice_code','title'], 'project':['topic']}.items()}
        config['practice'].update(start_date='2026-01-01',end_date='2026-01-02',supervisors=['Тестовый руководитель'])
        config['project'].update(path='project/example',assignment='project/example/assignment.md')
        path=self.root/'config/student.yaml'; path.write_text(yaml.safe_dump(config,allow_unicode=True),encoding='utf-8')
        first=prepare(self.root,path); self.assertEqual(first['status'],'ready')
        tool=self.root/'tools/check.py'; tool.write_text('version = 1',encoding='utf-8')
        with_tool=prepare(self.root,path)
        self.assertNotEqual(first['input_signature'],with_tool['input_signature'])
        self.assertIn('tools/check.py',with_tool['fingerprints'])
        tool.write_text('version = 2',encoding='utf-8')
        self.assertNotEqual(with_tool['input_signature'],prepare(self.root,path)['input_signature'])
        source.write_text('value = 2',encoding='utf-8')
        self.assertNotEqual(first['input_signature'],prepare(self.root,path)['input_signature'])
        config['practice']['end_date']='2025-12-31'; path.write_text(yaml.safe_dump(config),encoding='utf-8')
        self.assertEqual(prepare(self.root,path)['status'],'needs-input')

if __name__=='__main__':
    unittest.main()
