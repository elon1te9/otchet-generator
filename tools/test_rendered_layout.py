"""PDF regression tests and optional real Word smoke render.

Run unit tests normally; --word also exports a deliberately sparse synthetic
DOCX. It must pass structure and fail the 70% gate, without sentence-end errors.
No user application claims or imitated screenshots are generated.
"""
import argparse
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from rendered_layout import merged_height, page_occupancy, inspect_rendered, PT_PER_MM

ROOT = Path(__file__).resolve().parents[1]


def make_pdf(path, contents):
    """Minimal vector PDF fixtures using existing dependencies and Latin text."""
    writer = PdfWriter()
    for lines in contents:
        page = writer.add_blank_page(210 * PT_PER_MM, 297 * PT_PER_MM)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Times-Roman')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        commands = []
        for text, top in lines:
            escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            commands.append(f'BT /F1 14 Tf 1 0 0 1 90 {page.mediabox.height - top - 14} Tm ({escaped}) Tj ET')
        stream = DecodedStreamObject(); stream.set_data('\n'.join(commands).encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    with path.open('wb') as output:
        writer.write(output)


class RenderedLayoutTests(unittest.TestCase):
    def setUp(self):
        (ROOT/'temp').mkdir(exist_ok=True)
        self.workspace=tempfile.TemporaryDirectory(dir=ROOT/'temp')
        self.path=Path(self.workspace.name)/'fixture.pdf'
    def tearDown(self):
        self.workspace.cleanup()
    def test_whitespace_between_top_and_bottom_is_not_occupied(self):
        self.assertEqual(merged_height([(60,80),(700,720)],50,780),40)
    def test_overlapping_text_image_bands_are_not_double_counted(self):
        self.assertEqual(merged_height([(60,200),(80,90),(180,240)],50,780),180)
    def test_pending_blank_space_cannot_satisfy_occupancy(self):
        make_pdf(self.path,[[('SECTION',60),('Caption for pending figure',600),('Text.',630)]])
        pages,issues,_=inspect_rendered(self.path,[{'text':'SECTION','level':1}])
        self.assertLess(pages[0]['occupancy'],.15)
        self.assertIn('page-occupancy',{i['code'] for i in issues})
    def test_final_section_page_is_also_checked(self):
        full=[('SECTION',60)]+[(f'Line {i} continues',85+i*24) for i in range(26)]
        make_pdf(self.path,[[('TITLE',60)],[('CONTENTS',60)],full,[('End.',60)]])
        pages,issues,first=inspect_rendered(self.path,[{'text':'SECTION','level':1}])
        self.assertEqual(first,3)
        self.assertFalse(pages[0]['body']); self.assertFalse(pages[1]['body'])
        self.assertGreaterEqual(pages[2]['occupancy'],.70)
        self.assertFalse(pages[2]['sentence_end_preference_met'])
        self.assertEqual([i['code'] for i in issues],['page-occupancy'])
        self.assertIn('Page 4',issues[0]['message'])
    def test_footer_crossing_work_area_boundary_is_excluded(self):
        make_pdf(self.path,[[('SECTION',60),('Last body line.',90),('1',780)]])
        pages,_,_=inspect_rendered(self.path,[{'text':'SECTION','level':1}])
        self.assertEqual(pages[0]['last_line'],'Last body line.')
    def test_relaxed_sentence_end_does_not_allow_figure_caption_at_page_end(self):
        from unittest.mock import patch
        make_pdf(self.path,[[('SECTION',60),('Caption',90)]])
        lines=[{'text':'SECTION','top':62,'bottom':76,'size':14},
               {'text':'Рисунок 1.1 Подпись','top':92,'bottom':106,'size':14}]
        with patch('rendered_layout.content_lines',return_value=lines):
            _,issues,_=inspect_rendered(self.path,[{'text':'SECTION','level':1}])
        self.assertIn('pdf-figure-page-end',{i['code'] for i in issues})
        self.assertNotIn('sentence-page-end',{i['code'] for i in issues})
    def test_section_mid_page_and_orphan_heading_are_detected(self):
        make_pdf(self.path,[[('SECTION',60),('Content.',90),('SECOND',700)]])
        _,issues,_=inspect_rendered(self.path,[{'text':'SECTION','level':1},{'text':'SECOND','level':1}])
        self.assertTrue({'pdf-section-page-start','pdf-orphan-heading'} <= {i['code'] for i in issues})
    def test_multiple_subsections_can_share_page(self):
        lines=[('SECTION',60),('1.1 First',85),('Body.',110),('1.2 Second',160)]
        lines += [(f'Text {i} continues',185+i*24) for i in range(23)]
        make_pdf(self.path,[lines])
        headings=[{'text':'SECTION','level':1},{'text':'1.1 First','level':2},{'text':'1.2 Second','level':2}]
        pages,issues,_=inspect_rendered(self.path,headings)
        self.assertGreaterEqual(pages[0]['occupancy'],.70)
        self.assertEqual(issues,[])
    def test_image_height_counts(self):
        class Page:
            images=[{'x0':90,'x1':400,'top':200,'bottom':500}]
            def find_tables(self):
                return []
        metrics=page_occupancy(Page(),(80,60,560,780),[])
        self.assertAlmostEqual(metrics['occupancy'],300/720,places=3)
    def test_empty_table_height_does_not_count(self):
        class Table:
            bbox=(90,100,400,700)
        class Page:
            images=[]
            def find_tables(self):
                return [Table()]
        metrics=page_occupancy(Page(),(80,60,560,780),
                               [{'top':110,'bottom':122,'size':12}])
        self.assertLess(metrics['occupancy'],.03)


def smoke_word(output):
    from test_generator import synthetic_plan, synthetic_content, synthetic_manifest
    from build_report import build
    from check_revision_layout import check
    from check_subsection_conclusions import check as conclusions
    output=output.resolve()
    if not output.is_relative_to(ROOT/'temp'):
        raise ValueError('Regression artifacts must stay under temp/.')
    output.mkdir(parents=True,exist_ok=True)
    plan=synthetic_plan(); content=synthetic_content(); manifest=synthetic_manifest()
    # Extend one genuine test-input description to exercise automatic sentence
    # continuation. Repetition here is stress data, never report prose.
    content['blocks'][1]['text'] += ' '+('Синтетические слова для проверки автоматического переноса внутри предложения ' * 60).strip()+'.'
    for name,data in [('plan',plan),('content',content),('manifest',manifest)]:
        (output/f'{name}.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    title=Document(); title.add_paragraph('ПРОБНЫЙ ДОКУМЕНТ ГЕНЕРАТОРА')
    title.add_paragraph('Синтетический вход проверки верстки. Не отчет студента.')
    title.save(output/'title.docx')
    draft=build(output/'title.docx',content,plan,output/'draft.docx',ROOT)
    report=output/'report.docx'; pdf=output/'report.pdf'
    command=['powershell','-NoProfile','-ExecutionPolicy','Bypass','-File',str(ROOT/'tools/render_docx_word.ps1'),
             '-InputDocx',str(draft),'-OutputDocx',str(report),'-OutputPdf',str(pdf),'-UpdateFields']
    rendered=subprocess.run(command,capture_output=True,text=True,timeout=180)
    (output/'word.log').write_text(rendered.stdout+'\n'+rendered.stderr,encoding='utf-8')
    if rendered.returncode:
        raise RuntimeError('Word render failed; see '+str(output/'word.log'))
    structure=subprocess.run([sys.executable,str(ROOT/'tools/validate_report.py'),str(report),'--plan',str(output/'plan.json'),
                              '--json',str(output/'structure.json')],capture_output=True,text=True)
    (output/'structure.log').write_text(structure.stdout+'\n'+structure.stderr,encoding='utf-8')
    result=check(Document(report),pdf=pdf,plan=plan,manifest=manifest)
    (output/'layout.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    (output/'conclusions.json').write_text(json.dumps(conclusions(Document(report),plan),ensure_ascii=False,indent=2),encoding='utf-8')
    subprocess.run([sys.executable,str(ROOT/'tools/inspect_pdf.py'),str(pdf),'--json',str(output/'pdf.json'),
                    '--render-dir',str(output/'pages')],check=True)
    if len(list((output/'pages').glob('page-*.png'))) != len(result['pages']):
        raise AssertionError('PNG page count must match the latest PDF.')
    codes={i['code'] for i in result['issues']}
    if structure.returncode or codes != {'page-occupancy'}:
        raise AssertionError('Unexpected regression result: '+str(codes)+'; see structure/layout JSON.')
    if not any(p['body'] and not p['sentence_end_preference_met'] for p in result['pages']):
        raise AssertionError('Fixture must exercise a non-sentence page ending.')
    if not any(sum(h['level']==2 for h in p['headings']) >= 2 for p in result['pages']):
        raise AssertionError('At least two subsections must actually share a rendered page.')
    print(json.dumps({'regression':'pass','document_status':result['status'],
                      'expected_failure':'page-occupancy','pages':len(result['pages']),
                      'visual_review_required':True,'artifacts':str(output)},ensure_ascii=False,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--word',action='store_true')
    parser.add_argument('--output-dir',type=Path,default=ROOT/'temp/qa/render-regression')
    args,remaining=parser.parse_known_args()
    if args.word:
        smoke_word(args.output_dir)
    else:
        unittest.main(argv=[sys.argv[0],*remaining])
