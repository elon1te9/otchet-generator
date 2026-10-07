"""Regression tests for the user's report formatting rules; no application tests."""
import unittest
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_formatting import install_styles
from check_revision_layout import check

class LayoutGateTests(unittest.TestCase):
    def setUp(self):
        self.doc=Document(); install_styles(self.doc)
        self.heading=self.doc.add_paragraph('ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ','Practice Structural')
        self.doc.add_paragraph('Описание окна (Рисунок 2.1).','Practice Body')
        self.caption=self.doc.add_paragraph('Рисунок 2.1 Окно входа','Practice Caption')
        self.doc.add_paragraph('Завершенное предложение.','Practice Body')
        self.doc.add_paragraph('1 РАЗДЕЛ','Practice Section')
        self.subheading=self.doc.add_paragraph('1.1 Подраздел','Practice Subsection')
        self.doc.add_paragraph('Завершенное предложение.','Practice Body')
        self.table=self.doc.add_table(rows=1,cols=1)
        self.cell=self.table.cell(0,0).paragraphs[0]; self.cell.style='Practice Table'; self.cell.add_run('Данные')
    def codes(self): return {x['code'] for x in check(self.doc)['issues']}
    def test_valid_document(self): self.assertEqual(self.codes(),set())
    def test_bold_heading(self):
        self.heading.runs[0].bold=True; self.assertIn('bold-heading',self.codes())
    def test_structural_heading_alignment(self):
        self.heading.alignment=WD_ALIGN_PARAGRAPH.LEFT; self.assertIn('alignment',self.codes())
    def test_section_heading_indent(self):
        section=next(p for p in self.doc.paragraphs if p.text=='1 РАЗДЕЛ')
        section.paragraph_format.first_line_indent=Mm(0); self.assertIn('indent',self.codes())
    def test_table_spacing_and_indent(self):
        self.cell.paragraph_format.space_after=Pt(3)
        self.cell.paragraph_format.first_line_indent=Mm(5)
        self.assertTrue({'paragraph-spacing','indent'}<=self.codes())
    def test_caption_spacing(self):
        self.caption.paragraph_format.line_spacing=1; self.assertIn('line-spacing',self.codes())
    def test_long_dash(self):
        self.doc.paragraphs[1].add_run(' Текст — текст.'); self.assertIn('long-dash',self.codes())
    def test_reference_form(self):
        self.doc.paragraphs[1].text='Окно показано на рисунке 2.1.'
        self.assertIn('figure-reference',self.codes())
    def test_subheading_gap(self):
        before=self.doc.add_paragraph('','Practice Body')
        self.doc.add_paragraph('1.2 Следующий подраздел','Practice Subsection')
        self.assertNotIn('subsection-gap',self.codes())
        before.insert_paragraph_before(''); self.assertIn('subsection-gap',self.codes())
    def test_no_gap_after_section(self):
        self.subheading.insert_paragraph_before('').style='Practice Body'
        self.assertIn('section-subsection-gap',self.codes())
    def test_subheading_first_line_indent(self):
        self.subheading.paragraph_format.first_line_indent=Mm(0)
        self.assertIn('indent',self.codes())
    def test_table_caption_indent_and_alignment(self):
        caption=self.doc.add_paragraph('Таблица 3.1 - Тест-кейсы','Practice Table Caption')
        self.assertEqual(self.codes(),set())
        caption.paragraph_format.first_line_indent=Mm(0)
        caption.alignment=WD_ALIGN_PARAGRAPH.LEFT
        self.assertTrue({'indent','alignment'}<=self.codes())
    def test_ordinary_paragraph_pagination_markers(self):
        self.doc.paragraphs[1].paragraph_format.keep_together=True
        self.assertIn('pagination-marker',self.codes())
    def test_table_pagination_markers(self):
        self.cell.paragraph_format.keep_with_next=True
        self.assertIn('pagination-marker',self.codes())
    def test_guide_requires_illustrations(self):
        self.doc.add_paragraph('','Practice Body')
        self.doc.add_paragraph('2.4 Руководство пользователя','Practice Subsection')
        self.doc.add_paragraph('Порядок работы.','Practice Body')
        self.doc.add_paragraph('3 ТЕСТИРОВАНИЕ','Practice Section')
        self.assertIn('guide-figures',self.codes())
    def test_testing_requires_two_tables(self):
        self.doc.add_paragraph('3 ТЕСТИРОВАНИЕ','Practice Section')
        self.assertIn('testing-tables',self.codes())

if __name__=='__main__': unittest.main()
