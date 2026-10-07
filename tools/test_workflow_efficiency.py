"""Ensure context savings do not discard diagnostics, artifacts or checks."""
from contextlib import redirect_stdout
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from build_report import build
from prepare_report import compare_inputs, main as prepare_main
from report_cli import emit_result
from test_generator import synthetic_plan, synthetic_content, synthetic_manifest
from test_rendered_layout import make_pdf

ROOT = Path(__file__).resolve().parents[1]


def snapshot(files=None, status='ready'):
    hashes = {name: hashlib.sha256(value.encode()).hexdigest()
              for name, value in (files or {'project/main.py': 'one'}).items()}
    return {'status': status, 'issues': [], 'fingerprints': hashes,
            'input_signature': hashlib.sha256(json.dumps(hashes, sort_keys=True).encode()).hexdigest(),
            'execution_verified': False}


class CacheTests(unittest.TestCase):
    def test_only_complete_matching_snapshots_are_hits(self):
        current = snapshot()
        hit = compare_inputs(current, snapshot())
        self.assertEqual(hit['status'], 'match')
        self.assertTrue(hit['inputs_match'])
        self.assertFalse(hit['execution_verified'])
        for previous in (None, {}, [], snapshot(status='needs-input'),
                         dict(current, input_signature='tampered'),
                         dict(current, fingerprints={'a': 'bad'}),
                         dict(current, fingerprints={1: []})):
            with self.subTest(previous=previous):
                self.assertFalse(compare_inputs(current, previous)['inputs_match'])
        self.assertFalse(compare_inputs(snapshot(status='needs-input'), current)['inputs_match'])

    def test_modified_added_deleted_inputs_invalidate_hit(self):
        previous = snapshot({'project/main.py': 'one', 'project/old.py': 'old'})
        current = snapshot({'project/main.py': 'two', 'project/new.py': 'new'})
        delta = compare_inputs(current, previous)
        self.assertEqual(delta['status'], 'changed')
        self.assertEqual(delta['changed'], ['project/main.py'])
        self.assertEqual(delta['added'], ['project/new.py'])
        self.assertEqual(delta['removed'], ['project/old.py'])
        self.assertFalse(delta['inputs_match'])

    def test_previous_file_is_compared_before_overwrite(self):
        (ROOT/'temp').mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=ROOT/'temp') as directory:
            target = Path(directory)/'setup.json'
            target.write_text(json.dumps(snapshot()), encoding='utf-8')
            changed = snapshot({'project/main.py': 'two'})
            with patch('prepare_report.prepare', return_value=changed), \
                 patch.object(sys, 'argv', ['prepare_report', '--json', str(target), '--summary']), \
                 redirect_stdout(io.StringIO()) as console:
                self.assertEqual(prepare_main(), 0)
            saved = json.loads(target.read_text(encoding='utf-8'))
            self.assertEqual(saved['cache']['changed'], ['project/main.py'])
            self.assertEqual(saved['cache']['previous_input_signature'], snapshot()['input_signature'])
            self.assertEqual(json.loads(console.getvalue())['cache']['status'], 'changed')
            target.write_text('{broken', encoding='utf-8')
            with patch('prepare_report.prepare', return_value=snapshot()), \
                 patch.object(sys, 'argv', ['prepare_report', '--json', str(target), '--summary']), \
                 redirect_stdout(io.StringIO()):
                self.assertEqual(prepare_main(), 0)
            self.assertEqual(json.loads(target.read_text(encoding='utf-8'))['cache']['status'], 'invalid')


class OutputTests(unittest.TestCase):
    def setUp(self):
        (ROOT/'temp').mkdir(exist_ok=True)
        self.workspace = tempfile.TemporaryDirectory(dir=ROOT/'temp')
        self.root = Path(self.workspace.name)

    def tearDown(self):
        self.workspace.cleanup()

    def run_cli(self, name, *args):
        return subprocess.run([sys.executable, str(ROOT/'tools'/name), *map(str, args)],
                              capture_output=True, text=True, encoding='utf-8',
                              env=self.env(), timeout=60)

    @staticmethod
    def env():
        import os
        return dict(os.environ, PYTHONIOENCODING='utf-8')

    def test_full_artifact_and_all_diagnostics_are_preserved(self):
        result = {'status': 'fail', 'scope': 'structural-only',
                  'issues': [{'code': 'error', 'message': 'Ошибка'},
                             {'code': 'warning', 'severity': 'warning'}],
                  'pages': [{'text': 'Long page text '*100}]*30,
                  'pending_manual': ['Рисунок 1.1'], 'rendered_status': 'not-run',
                  'complete_verification': False, 'visual_review_required': True,
                  'evidence_review_required': True}
        target = self.root/'details.json'
        with redirect_stdout(io.StringIO()) as console:
            emit_result(result, target, summary=True)
        self.assertEqual(json.loads(target.read_text(encoding='utf-8')), result)
        summary = json.loads(console.getvalue())
        for key in ('status', 'scope', 'issues', 'pending_manual', 'rendered_status',
                    'complete_verification', 'visual_review_required', 'evidence_review_required'):
            self.assertEqual(summary[key], result[key])
        self.assertEqual(summary['page_count'], 30)
        self.assertLess(len(console.getvalue()), len(target.read_text())/10)
        with self.assertRaises(ValueError):
            emit_result(result, summary=True)

    def test_summary_requires_full_artifact_for_every_cli(self):
        for name in ('prepare_report.py', 'inspect_docx.py', 'inspect_pdf.py',
                     'validate_report.py', 'check_revision_layout.py', 'check_subsection_conclusions.py'):
            args = [] if name == 'prepare_report.py' else ['missing-file']
            if name in ('validate_report.py', 'check_revision_layout.py', 'check_subsection_conclusions.py'):
                args += ['--plan', 'missing-plan']
            result = self.run_cli(name, *args, '--summary')
            with self.subTest(name=name):
                self.assertEqual(result.returncode, 2)
                self.assertIn('--summary requires --json', result.stderr)

    def report_fixture(self):
        title = self.root/'title.docx'
        Document().save(title)
        report = self.root/'report.docx'
        build(title, synthetic_content(), synthetic_plan(), report, self.root)
        plan = self.root/'plan.json'
        plan.write_text(json.dumps(synthetic_plan(), ensure_ascii=False), encoding='utf-8')
        manifest = self.root/'manifest.json'
        manifest.write_text(json.dumps(synthetic_manifest(), ensure_ascii=False), encoding='utf-8')
        return report, plan, manifest

    def compare_cli_modes(self, name, args, extra=()):
        normal_path, summary_path = self.root/'normal.json', self.root/'summary.json'
        normal = self.run_cli(name, *args, '--json', normal_path, *extra)
        compact = self.run_cli(name, *args, '--json', summary_path, '--summary', *extra)
        self.assertEqual(normal.returncode, compact.returncode, compact.stderr)
        self.assertTrue(normal_path.exists(), normal.stderr)
        details = json.loads(normal_path.read_text(encoding='utf-8'))
        self.assertEqual(details, json.loads(summary_path.read_text(encoding='utf-8')))
        summary = json.loads(compact.stdout)
        if 'issues' in details:
            self.assertEqual(summary['issues'], details['issues'])
        if 'status' in details:
            self.assertEqual(summary['status'], details['status'])
        return details, summary, normal.returncode

    def test_report_commands_keep_identical_results_and_exit_codes(self):
        report, plan, manifest = self.report_fixture()
        for name, args in (
            ('inspect_docx.py', [report]),
            ('validate_report.py', [report, '--plan', plan]),
            ('check_revision_layout.py', [report, '--plan', plan, '--manifest', manifest]),
            ('check_subsection_conclusions.py', [report, '--plan', plan]),
        ):
            with self.subTest(name=name):
                details, summary, code = self.compare_cli_modes(name, args)
                if name == 'inspect_docx.py':
                    self.assertEqual(summary['counts'], details['counts'])
                if name == 'check_revision_layout.py':
                    self.assertEqual(code, 2)
                    self.assertEqual(summary['pending_manual'], details['pending_manual'])
                    self.assertFalse(summary['complete_verification'])
        doc = Document(report)
        doc.add_paragraph('TODO неподтвержденный текст')
        doc.save(report)
        details, summary, code = self.compare_cli_modes('validate_report.py', [report, '--plan', plan])
        self.assertEqual(code, 1)
        self.assertEqual(summary['status'], 'fail')

    def test_pdf_summary_keeps_text_and_every_rendered_page(self):
        pdf = self.root/'fixture.pdf'
        make_pdf(pdf, [[('First page', 250)], [('Second page', 250)]])
        outputs = []
        for mode in ('normal', 'summary'):
            folder = self.root/mode
            args = [pdf, '--json', folder/'pdf.json', '--text', folder/'pdf.txt',
                    '--render-dir', folder/'pages']
            if mode == 'summary':
                args += ['--summary']
            result = self.run_cli('inspect_pdf.py', *args)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(sorted(p.name for p in (folder/'pages').glob('*.png')),
                             ['page-1.png', 'page-2.png'])
            outputs.append(folder)
        for relative in ('pdf.json', 'pdf.txt', 'pages/page-1.png', 'pages/page-2.png'):
            self.assertEqual((outputs[0]/relative).read_bytes(), (outputs[1]/relative).read_bytes())


if __name__ == '__main__':
    unittest.main()
