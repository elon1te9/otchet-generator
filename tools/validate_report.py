#!/usr/bin/env python3
"""Deterministic structural checks for a completed practice-report DOCX."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from report_plan import load_plan, structure_issues, body_start, is_toc
from report_pagination import pagination_issues


FIGURE_RE = re.compile(r"^Рисунок\s+(\d+)\.(\d+)\s+\S")
TABLE_RE = re.compile(r"^Таблица\s+(\d+)\.(\d+)\s*-\s*\S")
PLACEHOLDER_RE = re.compile(r"(?:TODO|TBD|FIXME|<REQUIRED>|\[ВСТАВИТЬ|\{\{.+?\}\})", re.I)


def mm(length):
    return None if length is None else round(length / 36000, 2)


def near(actual, expected, tolerance=0.6):
    return actual is not None and abs(actual - expected) <= tolerance


def add_issue(issues, severity, code, message):
    issues.append({"severity": severity, "code": code, "message": message})


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--plan", required=True, type=Path, help="Current report structure from the selected example.")
    parser.add_argument("--json", type=Path)
    parser.add_argument("--protected-baseline", type=Path, help="Preserve the user's title section layout; verify it against this baseline.")
    args = parser.parse_args()

    doc = Document(args.input)
    plan = load_plan(args.plan)
    baseline = Document(args.protected_baseline) if args.protected_baseline else None
    issues = []
    begin = body_start(doc, plan)
    body = [p.text.strip() for p in doc.paragraphs[begin or 0:] if p.text.strip() and not is_toc(p)]
    joined = "\n".join(body)
    with zipfile.ZipFile(args.input) as zf:
        names = zf.namelist()
        document_xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        footer_xml = "".join(
            zf.read(name).decode("utf-8", errors="ignore")
            for name in names
            if name.startswith("word/footer") and name.endswith(".xml")
        )
        image_count = len([name for name in names if name.startswith("word/media/")])

    for section_index, section in enumerate(doc.sections, start=1):
        checks = {
            "page width": (mm(section.page_width), 210),
            "page height": (mm(section.page_height), 297),
            "top margin": (mm(section.top_margin), 20),
            "bottom margin": (mm(section.bottom_margin), 20),
            "left margin": (mm(section.left_margin), 30),
            "right margin": (mm(section.right_margin), 10),
        }
        for label, (actual, expected) in checks.items():
            if not near(actual, expected):
                attribute = {'page width':'page_width','page height':'page_height','top margin':'top_margin','bottom margin':'bottom_margin','left margin':'left_margin','right margin':'right_margin'}[label]
                preserved = baseline is not None and section_index == 1 and near(actual, mm(getattr(baseline.sections[0], attribute)))
                add_issue(issues, "info" if preserved else "error", "protected-section-layout" if preserved else "section-layout", f"Section {section_index}: {label} is {actual} mm; " + ("preserved from protected user title." if preserved else f"expected {expected} mm."))

    normal = doc.styles["Practice Body"] if "Practice Body" in doc.styles else doc.styles["Normal"]
    if normal.font.name not in (None, "Times New Roman"):
        add_issue(issues, "error", "body-font", f"Normal font is {normal.font.name!r}; expected Times New Roman.")
    if normal.font.size is not None and abs(normal.font.size.pt - 14) > 0.1:
        add_issue(issues, "error", "body-size", f"Normal size is {normal.font.size.pt} pt; expected 14 pt.")
    pf = normal.paragraph_format
    if pf.first_line_indent is not None and not near(mm(pf.first_line_indent), 12.5, 0.3):
        add_issue(issues, "error", "body-indent", f"Normal first-line indent is {mm(pf.first_line_indent)} mm; expected 12.5 mm.")
    if isinstance(pf.line_spacing, float) and abs(pf.line_spacing - 1.5) > 0.01:
        add_issue(issues, "error", "body-spacing", f"Normal line spacing is {pf.line_spacing}; expected 1.5.")

    for problem in structure_issues(doc, plan):
        add_issue(issues, "error", problem['code'], problem['message'])
    if begin is not None:
        for problem in pagination_issues(doc, begin):
            add_issue(issues, "error", problem['code'], problem['message'])
    if 'СОДЕРЖАНИЕ' not in document_xml:
        add_issue(issues, "error", "missing-toc-heading", "Missing СОДЕРЖАНИЕ heading.")

    figures = []
    tables = []
    for index, text in enumerate(body):
        fig = FIGURE_RE.match(text)
        if fig:
            figures.append((index, int(fig.group(1)), int(fig.group(2)), text))
            number = f"{fig.group(1)}.{fig.group(2)}"
            prior = " ".join(body[max(0, index - 4):index]).lower()
            if not re.search(rf"\(рисунок\s+{re.escape(number)}\)", prior, re.I):
                add_issue(issues, "warning", "figure-reference", f"No nearby reference before caption {text!r}.")
        tbl = TABLE_RE.match(text)
        if tbl:
            tables.append((index, int(tbl.group(1)), int(tbl.group(2)), text))

    for kind, records in (("figure", figures), ("table", tables)):
        per_section = {}
        for _index, section_no, item_no, caption in records:
            expected = per_section.get(section_no, 0) + 1
            if item_no != expected:
                add_issue(issues, "error", f"{kind}-sequence", f"Unexpected numbering in {caption!r}; expected {section_no}.{expected}.")
            per_section[section_no] = item_no

    if PLACEHOLDER_RE.search(joined):
        add_issue(issues, "error", "placeholder", "Unresolved placeholder text remains in the document.")

    if not re.search(r"\bTOC\b", document_xml):
        add_issue(issues, "error", "toc-field", "No Word TOC field found.")
    if not re.search(r"\bPAGE\b", footer_xml):
        add_issue(issues, "error", "page-field", "No PAGE field found in footers.")
    if image_count and not figures:
        add_issue(issues, "warning", "uncaptioned-images", f"Package contains {image_count} media files but no recognized figure captions.")

    result = {
        "file": str(args.input.resolve()),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "figure_caption_count": len(figures),
        "issues": issues,
        "status": "fail" if any(i["severity"] == "error" for i in issues) else "pass",
        "scope": "structural-only",
        "rendered_review_required": True,
    }
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(payload, encoding="utf-8")
    print(payload)
    raise SystemExit(1 if result["status"] == "fail" else 0)


if __name__ == "__main__":
    main()
